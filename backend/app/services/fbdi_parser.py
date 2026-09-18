"""
fbdi_parser.py — reads an FBDI .xlsm binary from app_fbdi_files,
extracts sheet/column metadata, and persists to app_fbdi_sheets + app_fbdi_columns.

Row 4 = column group labels (e.g. "Organization", "Customer Account")
Row 5 = actual column headers; headers prefixed with * are required
Data rows start at row 6.
Skips LOV and Instructions sheets (non-data).
"""

import io
import uuid

import openpyxl
import psycopg2

_SKIP_SHEETS = {"lov", "instructions"}
_PRIMARY_SHEET = "customers"  # sheet treated as is_primary=True


def parse_and_store(fbdi_file_id: str, file_content: bytes, conn) -> None:
    """Parse FBDI binary and upsert sheet + column metadata into DB."""
    wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)

    cursor = conn.cursor()
    # Wipe existing metadata for this file so re-uploads are clean
    cursor.execute("DELETE FROM app_fbdi_sheets WHERE fbdi_file_id = %s", (fbdi_file_id,))

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in _SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=4, max_row=5, values_only=True))
        if len(rows) < 2:
            continue

        group_row = rows[0]   # row 4
        header_row = rows[1]  # row 5

        # Count data rows (row 6 onward) — stop at first fully-empty row
        row_count = 0
        for data_row in ws.iter_rows(min_row=6, values_only=True):
            if all(c is None for c in data_row):
                break
            row_count += 1

        sheet_id = f"sht_{uuid.uuid4().hex[:8]}"
        is_primary = sheet_name.lower() == _PRIMARY_SHEET
        cursor.execute(
            "INSERT INTO app_fbdi_sheets (id, fbdi_file_id, sheet_name, row_count, is_primary) "
            "VALUES (%s, %s, %s, %s, %s)",
            (sheet_id, fbdi_file_id, sheet_name, row_count, is_primary),
        )

        for order, (group, header) in enumerate(zip(group_row, header_row)):
            if header is None:
                continue
            col_name = str(header).strip()
            if not col_name:
                continue
            is_required = col_name.startswith("*")
            clean_name = col_name.lstrip("*").strip()
            group_label = str(group).strip() if group else None

            cursor.execute(
                "INSERT INTO app_fbdi_columns "
                "(id, sheet_id, column_name, column_order, column_group, is_required) "
                "VALUES (%s, %s, %s, %s, %s, %s)",
                (f"col_{uuid.uuid4().hex[:8]}", sheet_id, clean_name, order, group_label, is_required),
            )

    wb.close()
    cursor.close()
