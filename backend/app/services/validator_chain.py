import os
import time
import io
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from typing import List, Optional, Dict, Any
from app.schemas.validation_schema import (
    ValidationChainReport,
    ValidationStepResult,
    FileDetectionResult
)
from app.services.file_detector import FileDetectorService, SUPPORTED_EXTENSIONS

class ValidationChainEngine:
    @staticmethod
    def reconcile_files(
        upload_dir: str,
        source_file_name: str,
        target_file_name: str,
        source_key: str,
        target_key: str,
        mappings: List[Dict[str, Any]],
        exclusions: List[Dict[str, Any]],
        rules: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        source = ValidationChainEngine._read_reconciliation_file(upload_dir, source_file_name)
        target = ValidationChainEngine._read_reconciliation_file(upload_dir, target_file_name)
        if source_key not in source.columns:
            raise ValueError(f'Source key "{source_key}" was not found.')
        if target_key not in target.columns:
            raise ValueError(f'Target key "{target_key}" was not found.')

        source = ValidationChainEngine._apply_exclusions(source, exclusions)
        target = ValidationChainEngine._apply_exclusions(target, exclusions)
        rule_failures = ValidationChainEngine._evaluate_rules(source, rules)
        source = source.copy()
        target = target.copy()
        source["__recon_key"] = source[source_key].map(ValidationChainEngine._normalise)
        target["__recon_key"] = target[target_key].map(ValidationChainEngine._normalise)
        source = source[source["__recon_key"] != ""]
        target = target[target["__recon_key"] != ""]
        source_by_key = {row["__recon_key"]: row for _, row in source.iterrows()}
        target_by_key = {row["__recon_key"]: row for _, row in target.iterrows()}
        comparisons = [m for m in mappings if m.get("enabled", True)] or [
            {"sourceColumn": source_key, "targetColumn": target_key, "transformType": "direct"}
        ]
        summary: Dict[str, Dict[str, int]] = {}
        discrepancies: List[Dict[str, Any]] = []
        for mapping in comparisons:
            field = str(mapping.get("targetColumn") or mapping.get("sourceColumn") or "value")
            summary[field] = {"matched": 0, "mismatched": 0, "missingSource": 0, "missingTarget": 0}

        matched = 0
        for key in sorted(set(source_by_key) | set(target_by_key)):
            source_row = source_by_key.get(key)
            target_row = target_by_key.get(key)
            if source_row is None:
                for mapping in comparisons:
                    summary[str(mapping.get("targetColumn") or mapping.get("sourceColumn") or "value")]["missingSource"] += 1
                discrepancies.append({"id": f"missing-source-{key}", "key": key, "field": "", "sourceValue": "", "targetValue": str(target_row.get(target_key, "")), "type": "missing_source"})
                continue
            if target_row is None:
                for mapping in comparisons:
                    summary[str(mapping.get("targetColumn") or mapping.get("sourceColumn") or "value")]["missingTarget"] += 1
                discrepancies.append({"id": f"missing-target-{key}", "key": key, "field": "", "sourceValue": str(source_row.get(source_key, "")), "targetValue": "", "type": "missing_target"})
                continue
            row_matched = True
            for mapping in comparisons:
                source_column = str(mapping.get("sourceColumn", ""))
                target_column = str(mapping.get("targetColumn", ""))
                field = target_column or source_column
                if source_column not in source_row or target_column not in target_row:
                    summary[field]["mismatched"] += 1
                    row_matched = False
                    continue
                source_value = ValidationChainEngine._transform_value(source_row[source_column], mapping)
                target_value = ValidationChainEngine._normalise(target_row[target_column])
                if source_value == target_value:
                    summary[field]["matched"] += 1
                else:
                    summary[field]["mismatched"] += 1
                    row_matched = False
                    discrepancies.append({"id": f"mismatch-{key}-{field}", "key": key, "field": field, "sourceValue": str(source_value), "targetValue": str(target_value), "type": "value_mismatch"})
            if row_matched:
                matched += 1

        total_source = len(source)
        total_target = len(target)
        return {
            "batchId": "backend", "runAt": pd.Timestamp.utcnow().isoformat(),
            "totalSource": total_source, "totalTarget": total_target, "matched": matched,
            "unmatchedSource": max(0, total_source - matched), "unmatchedTarget": max(0, total_target - matched),
            "matchRate": round((matched / total_source) * 100, 2) if total_source else 0,
            "discrepancies": discrepancies,
            "summary": [{"column": name, **counts} for name, counts in summary.items()],
            "rulesApplied": len([rule for rule in rules if rule.get("enabled", True)]),
            "ruleFailures": rule_failures,
            "exclusionsApplied": len([exclusion for exclusion in exclusions if exclusion.get("enabled", True)]),
        }

    @staticmethod
    def _read_reconciliation_file(upload_dir: str, file_name: str) -> pd.DataFrame:
        path = os.path.join(upload_dir, os.path.basename(file_name))
        if not os.path.isfile(path):
            raise ValueError(f'Uploaded file "{file_name}" was not found.')
        info = FileDetectorService.detect_file_and_sheets(path)
        frame = ValidationChainEngine._load_dataframe(path, info)
        if frame is None or not info.sheets:
            raise ValueError(f'No readable data was found in "{file_name}".')
        return frame.fillna("")

    @staticmethod
    def _normalise(value: Any) -> str:
        if pd.isna(value):
            return ""
        return str(value).strip().casefold()

    @staticmethod
    def _transform_value(value: Any, mapping: Dict[str, Any]) -> str:
        result = "" if pd.isna(value) else str(value).strip()
        transform = mapping.get("transformType", "direct")
        if transform == "trim":
            return result.casefold()
        if transform == "upper":
            return result.upper().casefold()
        if transform == "lower":
            return result.lower().casefold()
        return result.casefold()

    @staticmethod
    def _apply_exclusions(frame: pd.DataFrame, exclusions: List[Dict[str, Any]]) -> pd.DataFrame:
        result = frame
        for exclusion in exclusions:
            column = exclusion.get("column")
            if not exclusion.get("enabled", True) or column not in result.columns:
                continue
            values = result[column].fillna("").astype(str)
            operand = str(exclusion.get("value", ""))
            operator = exclusion.get("operator", "equals")
            if operator == "equals":
                mask = values.str.casefold() == operand.casefold()
            elif operator == "contains":
                mask = values.str.contains(operand, case=False, na=False)
            elif operator == "startsWith":
                mask = values.str.startswith(operand, na=False)
            elif operator == "endsWith":
                mask = values.str.endswith(operand, na=False)
            elif operator == "isNull":
                mask = values.eq("")
            elif operator == "regex":
                mask = values.str.contains(operand, case=False, na=False, regex=True)
            elif operator in {"greaterThan", "lessThan"}:
                numeric = pd.to_numeric(values, errors="coerce")
                threshold = pd.to_numeric(operand, errors="coerce")
                mask = numeric > threshold if operator == "greaterThan" else numeric < threshold
            else:
                mask = pd.Series(False, index=result.index)
            result = result.loc[~mask]
        return result

    @staticmethod
    def _evaluate_rules(frame: pd.DataFrame, rules: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        failures: List[Dict[str, Any]] = []
        for rule in rules:
            if not rule.get("enabled", True) or rule.get("column") not in frame.columns:
                continue
            column = str(rule["column"])
            values = frame[column].fillna("").astype(str)
            config = rule.get("config") or {}
            rule_type = rule.get("type")
            if rule_type == "regex":
                pattern = str(config.get("pattern", ""))
                passed = values.str.contains(pattern, regex=True, na=False) if pattern else pd.Series(True, index=frame.index)
            elif rule_type == "range":
                numeric = pd.to_numeric(values, errors="coerce")
                minimum = config.get("min")
                maximum = config.get("max")
                passed = numeric.notna()
                if minimum is not None: passed &= numeric >= float(minimum)
                if maximum is not None: passed &= numeric <= float(maximum)
            elif rule_type == "lookup":
                allowed = {str(item).casefold() for item in config.get("values", [])}
                passed = values.str.casefold().isin(allowed)
            elif rule_type == "format":
                passed = values.str.strip().ne("")
            else:
                passed = pd.Series(True, index=frame.index)
            failed = int((~passed).sum())
            if failed:
                failures.append({"rule": rule.get("name", rule_type), "column": column, "severity": rule.get("severity", "warning"), "failed": failed, "total": len(frame)})
        return failures

    @staticmethod
    def build_report_workbook(report: Dict[str, Any], batch_name: str, source_file: str, target_file: str) -> bytes:
        template_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "templates", "AiretechCustomerRecon.xlsx"))
        workbook = load_workbook(template_path) if os.path.isfile(template_path) else Workbook()
        required_sheets = ["Recon", "Diff Customers", "Diff Contacts", "Diff Sites", "Fusion Data", "FusionContacts"]
        if workbook.active and workbook.active.title not in required_sheets and len(workbook.sheetnames) == 1:
            workbook.remove(workbook.active)
        for name in required_sheets:
            if name not in workbook.sheetnames:
                workbook.create_sheet(name)
        template_headers = {
            name: [cell.value for cell in workbook[name][1]]
            for name in required_sheets
        }
        for sheet in workbook.worksheets:
            for table_name in list(sheet.tables):
                del sheet.tables[table_name]
            if sheet.max_row:
                sheet.delete_rows(1, sheet.max_row)
        sheets = {name: workbook[name] for name in required_sheets}
        recon = sheets["Recon"]
        recon_rows = [
            [batch_name], [],
            ["Source File records", "Total_Count", None, "FBDI File Records", "TOTAL_COUNT"],
            ["Total Source Records", report["totalSource"], None, "Target Records", report["totalTarget"]],
            ["Matched Records", report["matched"], None, "Match Rate", report["matchRate"] / 100],
            ["Unmatched Source", report["unmatchedSource"], None, "Unmatched Target", report["unmatchedTarget"]],
            [], ["File", "File Path", None, None, None],
            ["Source File:", source_file, None, None, None],
            ["Target File:", target_file, None, None, None], [],
            ["Field", "Matched", "Mismatched", "Missing Source", "Missing Target"],
        ]
        for row in recon_rows:
            recon.append(row)
        for item in report.get("summary", []):
            recon.append([item["column"], item["matched"], item["mismatched"], item["missingSource"], item["missingTarget"]])

        customer_sheet = sheets["Diff Customers"]
        customer_sheet.append(["Customer Name", "Reason", "Status"])
        contact_sheet = sheets["Diff Contacts"]
        contact_sheet.append(["Account Name", "First Name", "Last Name", "Phone1", "Phone", "Fax", "Email", "Reason"])
        site_sheet = sheets["Diff Sites"]
        site_sheet.append(["Site Key", "Reason", "Status", "Source/Target"])
        for item in report.get("discrepancies", []):
            reason = str(item.get("type", "")).replace("_", " ").upper()
            status = "MISMATCH" if item.get("type") == "value_mismatch" else "MISSING"
            customer_sheet.append([item.get("key", ""), reason, status])
            contact_sheet.append([item.get("key", ""), "", "", "", "", "", "", reason])
            site_sheet.append([item.get("key", ""), reason, status, "Source" if item.get("type") == "missing_target" else "Target"])

        fusion_data = sheets["Fusion Data"]
        fusion_data.append(template_headers["Fusion Data"])
        fusion_contacts = sheets["FusionContacts"]
        fusion_contacts.append(template_headers["FusionContacts"])
        for item in report.get("discrepancies", []):
            fusion_data.append([item.get("key", ""), item.get("field", ""), item.get("sourceValue", ""), item.get("targetValue", ""), item.get("type", "")])
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for sheet in sheets.values():
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for column in range(1, sheet.max_column + 1):
                letter = get_column_letter(column)
                width = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 100) + 1)), default=14)
                sheet.column_dimensions[letter].width = min(48, max(14, width + 2))
            sheet.freeze_panes = "A2"
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def execute_validation_chain(
        folder_path: str,
        target_file_name: Optional[str] = None,
        file_type: str = "SOURCE",
        required_columns: Optional[List[str]] = None,
        primary_key_column: Optional[str] = None
    ) -> ValidationChainReport:
        start_time = time.time()
        steps: List[ValidationStepResult] = []
        required_columns = required_columns or []

        # -------------------------------------------------------------
        # STEP 1: Folder Validation
        # -------------------------------------------------------------
        folder_exists = os.path.exists(folder_path) and os.path.isdir(folder_path)
        if not folder_exists:
            steps.append(ValidationStepResult(
                step_number=1,
                step_name="Folder Validation",
                status="FAIL",
                message=f"Folder path '{folder_path}' does not exist or is not a directory.",
                details={"folder_path": folder_path}
            ))
            return ValidationChainReport(
                folder_path=folder_path,
                folder_exists=False,
                has_supported_file=False,
                overall_status="FAIL",
                validation_steps=steps,
                execution_time_ms=round((time.time() - start_time) * 1000, 2)
            )
        
        steps.append(ValidationStepResult(
            step_number=1,
            step_name="Folder Validation",
            status="PASS",
            message=f"Folder path '{folder_path}' exists.",
            details={"folder_path": folder_path}
        ))

        # -------------------------------------------------------------
        # STEP 2: File Validation (Does at least one supported file exist?)
        # -------------------------------------------------------------
        folder_files = os.listdir(folder_path)
        supported_files = [
            f for f in folder_files 
            if os.path.isfile(os.path.join(folder_path, f)) and FileDetectorService.is_supported_file(f)
        ]

        if target_file_name:
            supported_files = [f for f in supported_files if f == target_file_name]

        if not supported_files:
            steps.append(ValidationStepResult(
                step_number=2,
                step_name="File Format Validation",
                status="FAIL",
                message=f"No supported file found in folder. Supported formats: {sorted(list(SUPPORTED_EXTENSIONS))}",
                details={"found_files": folder_files}
            ))
            return ValidationChainReport(
                folder_path=folder_path,
                folder_exists=True,
                has_supported_file=False,
                overall_status="FAIL",
                validation_steps=steps,
                execution_time_ms=round((time.time() - start_time) * 1000, 2)
            )

        # Select target file: either specified file or first supported file
        selected_file_name = target_file_name if target_file_name in supported_files else supported_files[0]
        selected_file_path = os.path.join(folder_path, selected_file_name)

        steps.append(ValidationStepResult(
            step_number=2,
            step_name="File Format Validation",
            status="PASS",
            message=f"Found supported file '{selected_file_name}'. Total supported files: {len(supported_files)}",
            details={"selected_file": selected_file_name, "all_supported_files": supported_files}
        ))

        # Detect File & Sheet Meta
        file_info: FileDetectionResult = FileDetectorService.detect_file_and_sheets(selected_file_path, file_type=file_type)

        # -------------------------------------------------------------
        # STEP 3: File Size > 0
        # -------------------------------------------------------------
        if file_info.file_size_bytes == 0:
            steps.append(ValidationStepResult(
                step_number=3,
                step_name="File Size Validation",
                status="FAIL",
                message=f"File '{selected_file_name}' is empty (0 bytes).",
                details={"file_size_bytes": 0}
            ))
            return ValidationChainReport(
                folder_path=folder_path,
                folder_exists=True,
                has_supported_file=True,
                overall_status="FAIL",
                file_info=file_info,
                validation_steps=steps,
                execution_time_ms=round((time.time() - start_time) * 1000, 2)
            )

        steps.append(ValidationStepResult(
            step_number=3,
            step_name="File Size Validation",
            status="PASS",
            message=f"File size is valid: {file_info.file_size_bytes} bytes.",
            details={"file_size_bytes": file_info.file_size_bytes}
        ))

        # Read Tabular Data for Data Validation Execution Chain
        df = ValidationChainEngine._load_dataframe(selected_file_path, file_info)
        total_records = len(df) if df is not None else 0

        # -------------------------------------------------------------
        # STEP 4: Record Count > 0
        # -------------------------------------------------------------
        if total_records == 0:
            steps.append(ValidationStepResult(
                step_number=4,
                step_name="Record Count Validation",
                status="FAIL",
                message=f"File '{selected_file_name}' contains 0 records.",
                details={"record_count": 0}
            ))
            return ValidationChainReport(
                folder_path=folder_path,
                folder_exists=True,
                has_supported_file=True,
                overall_status="FAIL",
                file_info=file_info,
                validation_steps=steps,
                execution_time_ms=round((time.time() - start_time) * 1000, 2)
            )

        steps.append(ValidationStepResult(
            step_number=4,
            step_name="Record Count Validation",
            status="PASS",
            message=f"Record count is valid: {total_records} rows.",
            details={"record_count": total_records}
        ))

        # -------------------------------------------------------------
        # STEP 5: Required Columns Available Check
        # -------------------------------------------------------------
        detected_columns = [str(c) for c in df.columns.tolist()] if df is not None else []
        missing_cols = [col for col in required_columns if col not in detected_columns]

        if missing_cols:
            steps.append(ValidationStepResult(
                step_number=5,
                step_name="Required Columns Check",
                status="FAIL",
                message=f"Missing required columns: {missing_cols}",
                details={"required_columns": required_columns, "missing_columns": missing_cols, "detected_columns": detected_columns}
            ))
        else:
            steps.append(ValidationStepResult(
                step_number=5,
                step_name="Required Columns Check",
                status="PASS",
                message="All required columns are present.",
                details={"required_columns": required_columns, "detected_columns": detected_columns}
            ))

        # -------------------------------------------------------------
        # STEP 6: Duplicate Records Check
        # -------------------------------------------------------------
        duplicate_count = 0
        duplicate_details = {}

        if df is not None and not df.empty:
            if primary_key_column and primary_key_column in df.columns:
                pk_dups = df[df.duplicated(subset=[primary_key_column], keep=False)]
                duplicate_count = len(pk_dups)
                duplicate_details = {"check_type": "PRIMARY_KEY", "primary_key": primary_key_column, "duplicate_rows_count": duplicate_count}
            else:
                full_dups = df[df.duplicated(keep=False)]
                duplicate_count = len(full_dups)
                duplicate_details = {"check_type": "FULL_ROW", "duplicate_rows_count": duplicate_count}

        if duplicate_count > 0:
            steps.append(ValidationStepResult(
                step_number=6,
                step_name="Duplicate Records Check",
                status="WARNING",
                message=f"Found {duplicate_count} duplicate records.",
                details=duplicate_details
            ))
        else:
            steps.append(ValidationStepResult(
                step_number=6,
                step_name="Duplicate Records Check",
                status="PASS",
                message="No duplicate records detected.",
                details=duplicate_details
            ))

        # -------------------------------------------------------------
        # STEP 7: Null Values Check
        # -------------------------------------------------------------
        null_counts: Dict[str, int] = {}
        total_nulls = 0

        if df is not None and not df.empty:
            null_series = df.isnull().sum()
            null_counts = {col: int(cnt) for col, cnt in null_series.items() if cnt > 0}
            total_nulls = sum(null_counts.values())

        if total_nulls > 0:
            steps.append(ValidationStepResult(
                step_number=7,
                step_name="Null Values Check",
                status="WARNING",
                message=f"Detected {total_nulls} null/empty fields across columns.",
                details={"null_counts_per_column": null_counts, "total_nulls": total_nulls}
            ))
        else:
            steps.append(ValidationStepResult(
                step_number=7,
                step_name="Null Values Check",
                status="PASS",
                message="No null values detected.",
                details={"total_nulls": 0}
            ))

        # Overall Status Calculation
        has_fail = any(s.status == "FAIL" for s in steps)
        has_warning = any(s.status == "WARNING" for s in steps)
        overall_status = "FAIL" if has_fail else ("WARNING" if has_warning else "PASS")

        return ValidationChainReport(
            folder_path=folder_path,
            folder_exists=True,
            has_supported_file=True,
            overall_status=overall_status,
            file_info=file_info,
            validation_steps=steps,
            record_count=total_records,
            duplicate_records_count=duplicate_count,
            null_values_summary=null_counts,
            missing_required_columns=missing_cols,
            execution_time_ms=round((time.time() - start_time) * 1000, 2)
        )

    @staticmethod
    def _load_dataframe(file_path: str, file_info: FileDetectionResult) -> Optional[pd.DataFrame]:
        ext = file_info.file_extension.lower()
        try:
            if ext in ['.xlsx', '.xls']:
                engine = 'openpyxl' if ext == '.xlsx' else 'xlrd'
                with pd.ExcelFile(file_path, engine=engine) as excel_file:
                    return FileDetectorService.load_excel_sheet(excel_file, excel_file.sheet_names[0])
            elif ext in ['.csv', '.txt', '.dat']:
                delimiter = file_info.delimiter or ','
                encoding = file_info.encoding or 'utf-8'
                return pd.read_csv(file_path, sep=delimiter, encoding=encoding, low_memory=False)
            elif ext == '.json':
                return pd.read_json(file_path)
            elif ext == '.xml':
                return pd.read_xml(file_path)
        except Exception:
            return None
        return None
